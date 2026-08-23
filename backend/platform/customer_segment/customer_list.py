from __future__ import annotations

import base64
import binascii
import hashlib
import io
from datetime import date, datetime, time, timezone
from typing import Any

from openpyxl import load_workbook

from backend.platform.ingestion.workbook import MAX_WORKBOOK_BYTES, _validate_xlsx_archive


MAX_CUSTOMER_IDS = 20_000
MAX_CUSTOMER_LIST_ROWS = 50_000
MAX_CUSTOMER_ID_LENGTH = 128


def preview_customer_list(file_name: str, content_base64: str) -> dict[str, Any]:
    content = _decode_customer_list(file_name, content_base64)
    parsed = _parse_customer_ids(content)
    return {
        "file_name": _safe_file_name(file_name),
        "content_hash": hashlib.sha256(content).hexdigest(),
        "customer_count": len(parsed["customer_ids"]),
        "duplicate_count": parsed["duplicate_count"],
        "blank_count": parsed["blank_count"],
        "other_columns_ignored": parsed["other_columns_ignored"],
        "valid": True,
    }


def confirm_customer_list(
    services: Any,
    *,
    tenant_id: str,
    user_id: str,
    file_name: str,
    content_base64: str,
    expected_content_hash: str,
) -> dict[str, Any]:
    content = _decode_customer_list(file_name, content_base64)
    content_hash = hashlib.sha256(content).hexdigest()
    if not expected_content_hash or content_hash != str(expected_content_hash).strip().lower():
        raise ValueError("customer_segment_preview_stale")
    parsed = _parse_customer_ids(content)
    normalized = ("\n".join(parsed["customer_ids"]) + "\n").encode("utf-8")
    stored = services.data_acquisition_service.object_store.put(tenant_id, normalized, ".txt")
    artifact = services.data_acquisition_service.store.create_artifact(
        tenant_id,
        object_uri=stored.object_uri,
        content_hash=stored.content_hash,
        content_type="text/plain; charset=utf-8",
        size_bytes=stored.size_bytes,
        status="active",
        created_by=user_id,
        artifact_type="other",
    )
    confirmed_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    metadata = {
        "artifactId": str(artifact["artifact_id"]),
        "contentHash": str(artifact["content_hash"]),
        "sourceContentHash": content_hash,
        "fileName": _safe_file_name(file_name),
        "customerCount": len(parsed["customer_ids"]),
        "duplicateCount": parsed["duplicate_count"],
        "blankCount": parsed["blank_count"],
        "confirmedAt": confirmed_at,
        "ownerUserId": user_id,
    }
    result = services.application_store.run_action(
        tenant_id,
        "customer_segment_analysis",
        "set_customer_segment_list",
        payload={"customerList": metadata},
        actor_user_id=user_id,
    )
    return {
        "customer_list": metadata,
        "module": result["module"],
    }


def customer_ids_for_user(services: Any, *, tenant_id: str, user_id: str) -> tuple[dict[str, Any], list[str]]:
    module = services.application_store.get_module(
        tenant_id,
        "customer_segment_analysis",
        actor_user_id=user_id,
    )
    state = module.get("state") if isinstance(module, dict) else {}
    metadata = state.get("customerSegmentList") if isinstance(state, dict) else None
    if not isinstance(metadata, dict) or not str(metadata.get("artifactId") or ""):
        raise PermissionError("customer_segment_list_required")
    if str(metadata.get("ownerUserId") or "") != user_id:
        raise PermissionError("customer_segment_list_required")
    artifact, content = services.data_acquisition_service.get_artifact_content(
        tenant_id,
        str(metadata["artifactId"]),
    )
    if str(artifact.get("content_hash") or "") != str(metadata.get("contentHash") or ""):
        raise PermissionError("customer_segment_list_content_changed")
    try:
        customer_ids = [line.strip() for line in content.decode("utf-8").splitlines() if line.strip()]
    except UnicodeDecodeError as exc:
        raise PermissionError("customer_segment_list_content_invalid") from exc
    if not customer_ids or len(customer_ids) > MAX_CUSTOMER_IDS or len(customer_ids) != len(set(customer_ids)):
        raise PermissionError("customer_segment_list_content_invalid")
    if int(metadata.get("customerCount") or 0) != len(customer_ids):
        raise PermissionError("customer_segment_list_content_changed")
    return dict(metadata), customer_ids


def _decode_customer_list(file_name: str, content_base64: str) -> bytes:
    normalized_name = _safe_file_name(file_name)
    if not normalized_name.lower().endswith(".xlsx"):
        raise ValueError("customer_segment_list_requires_xlsx")
    try:
        content = base64.b64decode(str(content_base64 or ""), validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("invalid_content_base64") from exc
    if not content or len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError("customer_segment_list_size_invalid")
    _validate_xlsx_archive(content)
    return content


def _parse_customer_ids(content: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("customer_segment_list_workbook_invalid") from exc
    try:
        if not workbook.worksheets:
            raise ValueError("customer_segment_list_empty")
        worksheet = workbook.worksheets[0]
        customer_ids: list[str] = []
        seen: set[str] = set()
        duplicate_count = 0
        blank_count = 0
        other_columns_ignored = False
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number > MAX_CUSTOMER_LIST_ROWS:
                raise ValueError("customer_segment_list_row_limit_exceeded")
            value = row[0] if row else None
            if any(cell is not None and cell != "" for cell in row[1:]):
                other_columns_ignored = True
            customer_id = _customer_id_text(value, row_number)
            if not customer_id:
                blank_count += 1
                continue
            if customer_id in seen:
                duplicate_count += 1
                continue
            seen.add(customer_id)
            customer_ids.append(customer_id)
            if len(customer_ids) > MAX_CUSTOMER_IDS:
                raise ValueError("customer_segment_list_row_limit_exceeded")
        if not customer_ids:
            raise ValueError("customer_segment_list_empty")
        return {
            "customer_ids": customer_ids,
            "duplicate_count": duplicate_count,
            "blank_count": blank_count,
            "other_columns_ignored": other_columns_ignored,
        }
    finally:
        workbook.close()


def _customer_id_text(value: Any, row_number: int) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        normalized = value.strip()
    elif isinstance(value, bool) or isinstance(value, (date, datetime, time)):
        raise ValueError(f"customer_segment_list_value_invalid:{row_number}")
    elif isinstance(value, int):
        normalized = str(value)
    elif isinstance(value, float) and value.is_integer():
        normalized = str(int(value))
    else:
        raise ValueError(f"customer_segment_list_value_invalid:{row_number}")
    if len(normalized) > MAX_CUSTOMER_ID_LENGTH or any(ord(character) < 32 for character in normalized):
        raise ValueError(f"customer_segment_list_value_invalid:{row_number}")
    return normalized


def _safe_file_name(value: str) -> str:
    normalized = str(value or "").strip()
    if not normalized or len(normalized) > 240 or "/" in normalized or "\\" in normalized:
        raise ValueError("customer_segment_list_file_name_invalid")
    return normalized

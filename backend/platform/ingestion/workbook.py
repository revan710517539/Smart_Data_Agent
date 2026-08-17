from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook


MAX_WORKBOOK_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000
MAX_SHEETS = 30
MAX_ROWS_PER_SHEET = 50_000
MAX_TOTAL_ROWS = 100_000
MAX_COLUMNS = 300


@dataclass(frozen=True)
class WorkbookSheetCSV:
    name: str
    csv_bytes: bytes
    row_count: int
    column_count: int


def parse_static_workbook(file_name: str, content: bytes) -> list[WorkbookSheetCSV]:
    normalized_name = str(file_name or "").strip()
    if not normalized_name.lower().endswith(".xlsx"):
        raise ValueError("static_workbook_requires_xlsx")
    if not content or len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError("invalid_static_workbook_size")
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("invalid_static_workbook") from exc
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ValueError("static_workbook_sheet_limit_exceeded")
        sheets: list[WorkbookSheetCSV] = []
        total_rows = 0
        used_names: set[str] = set()
        for worksheet in workbook.worksheets:
            sheet = _worksheet_to_csv(worksheet.title, worksheet.iter_rows(values_only=True), used_names)
            if sheet is None:
                continue
            total_rows += sheet.row_count
            if total_rows > MAX_TOTAL_ROWS:
                raise ValueError("static_workbook_total_row_limit_exceeded")
            sheets.append(sheet)
        if not sheets:
            raise ValueError("static_workbook_has_no_non_empty_sheet")
        return sheets
    finally:
        workbook.close()


def safe_workbook_stem(file_name: str) -> str:
    stem = str(file_name or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    stem = stem.rsplit(".", 1)[0].strip()
    normalized = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", stem).strip(" ._")
    return (normalized or "静态工作簿")[:100]


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("static_workbook_archive_entry_limit_exceeded")
            total_size = 0
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ValueError("static_workbook_archive_path_invalid")
                total_size += int(entry.file_size)
                if total_size > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("static_workbook_uncompressed_limit_exceeded")
                if entry.compress_size > 0 and entry.file_size / entry.compress_size > 200:
                    raise ValueError("static_workbook_compression_ratio_invalid")
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset({entry.filename for entry in entries}):
                raise ValueError("invalid_static_workbook")
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid_static_workbook") from exc


def _worksheet_to_csv(
    original_name: str,
    source_rows: Any,
    used_names: set[str],
) -> WorkbookSheetCSV | None:
    rows: list[list[str]] = []
    started = False
    max_width = 0
    for raw_row in source_rows:
        values = [_cell_text(value) for value in raw_row]
        while values and not values[-1]:
            values.pop()
        if not started and not any(values):
            continue
        started = True
        if not any(values):
            continue
        max_width = max(max_width, len(values))
        if max_width > MAX_COLUMNS:
            raise ValueError("static_workbook_column_limit_exceeded")
        rows.append(values)
        if len(rows) - 1 > MAX_ROWS_PER_SHEET:
            raise ValueError("static_workbook_sheet_row_limit_exceeded")
    if not rows:
        return None
    headers = _normalized_headers(rows[0], max_width)
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(headers)
    for row in rows[1:]:
        writer.writerow([*(row[:max_width]), *([""] * max(0, max_width - len(row)))])
    name = _unique_sheet_name(original_name, used_names)
    return WorkbookSheetCSV(
        name=name,
        csv_bytes=output.getvalue().encode("utf-8-sig"),
        row_count=max(0, len(rows) - 1),
        column_count=max_width,
    )


def _normalized_headers(values: list[str], width: int) -> list[str]:
    headers: list[str] = []
    used: set[str] = set()
    for index in range(width):
        base = (values[index] if index < len(values) else "").strip() or f"字段{index + 1}"
        candidate = base[:120]
        suffix = 2
        while candidate in used:
            candidate = f"{base[:110]}_{suffix}"
            suffix += 1
        used.add(candidate)
        headers.append(candidate)
    return headers


def _unique_sheet_name(value: str, used_names: set[str]) -> str:
    base = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", str(value or "")).strip(" ._")[:80] or "Sheet"
    candidate = base
    suffix = 2
    while candidate.casefold() in used_names:
        candidate = f"{base[:72]}_{suffix}"
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

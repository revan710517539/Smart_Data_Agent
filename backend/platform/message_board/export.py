from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

EXPORT_HEADERS = (
    "留言时间",
    "留言人",
    "留言内容",
    "留言页面",
    "引用内容",
    "机构",
    "附件",
    "来源",
    "更新时间",
    "追加内容",
)
_ATTACHMENT_COLUMN = 6
_MAX_IMAGE_WIDTH = 160
_MAX_IMAGE_HEIGHT = 96
_CHINA = timezone(timedelta(hours=8))


class _EmbeddedImage(Image):
    def __init__(self, data: bytes, width: int, height: int, fmt: str) -> None:
        self.ref = data
        self.width = width
        self.height = height
        self.format = fmt
        self.anchor = "A1"

    def _data(self) -> bytes:
        return self.ref


def build_adopted_workbook(messages: list[dict[str, Any]], images: dict[str, list[bytes]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "已采纳留言"
    header_font = Font(bold=True, color="1D1D1F", size=11)
    header_fill = PatternFill("solid", fgColor="FAFBFC")
    wrap = Alignment(vertical="top", wrap_text=True)
    for index, title in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(1, index, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    widths = (20, 16, 42, 18, 28, 16, 32, 20, 20, 36)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "A2"

    for row_index, message in enumerate(messages, start=2):
        quote = message.get("quote_context") if isinstance(message.get("quote_context"), dict) else {}
        quote_text = str((quote or {}).get("selected_text") or "")
        page_key = str(message.get("page_key") or "")
        source = "登录页数据使用调查" if page_key == "login-survey" else str(message.get("page_title") or "")
        blobs = [item for item in images.get(str(message.get("message_id") or ""), []) if item]
        values = (
            _format_datetime(message.get("created_at")),
            str(message.get("author_name") or ""),
            str(message.get("content") or ""),
            str(message.get("page_title") or ""),
            quote_text,
            str(message.get("tenant_id") or "").replace("tenant:", ""),
            "" if blobs else f"{len(message.get('attachment_ids') or [])} 张截图",
            source,
            _format_datetime(message.get("updated_at")),
            str(message.get("append_content") or ""),
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = wrap
        image_height = _place_images(sheet, row_index, blobs)
        sheet.row_dimensions[row_index].height = max(28, image_height * 0.75 + 10)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _place_images(sheet: Any, row_index: int, blobs: list[bytes]) -> int:
    offset_x = 4
    tallest = 0
    for blob in blobs:
        width, height = _fit(*_image_size(blob))
        image = _EmbeddedImage(blob, width, height, _image_format(blob))
        image.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=_ATTACHMENT_COLUMN,
                colOff=pixels_to_EMU(offset_x),
                row=row_index - 1,
                rowOff=pixels_to_EMU(4),
            ),
            ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
        )
        sheet.add_image(image)
        offset_x += width + 8
        tallest = max(tallest, height)
    return tallest


def _fit(width: int, height: int) -> tuple[int, int]:
    safe_width = max(1, int(width or 1))
    safe_height = max(1, int(height or 1))
    scale = min(_MAX_IMAGE_WIDTH / safe_width, _MAX_IMAGE_HEIGHT / safe_height, 1.0)
    return max(1, int(safe_width * scale)), max(1, int(safe_height * scale))


def _image_format(data: bytes) -> str:
    if data.startswith(b"\x89PNG"):
        return "png"
    if data[:2] == b"\xff\xd8":
        return "jpeg"
    if data[:6] in {b"GIF87a", b"GIF89a"}:
        return "gif"
    return "png"


def _image_size(data: bytes) -> tuple[int, int]:
    if data.startswith(b"\x89PNG\r\n\x1a\n") and len(data) >= 24:
        return int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")
    if data[:2] == b"\xff\xd8":
        offset = 2
        while offset + 9 < len(data):
            if data[offset] != 0xFF:
                break
            marker = data[offset + 1]
            if marker in {0xD8, 0xD9}:
                offset += 2
                continue
            block = int.from_bytes(data[offset + 2:offset + 4], "big")
            if marker in {0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF}:
                return int.from_bytes(data[offset + 7:offset + 9], "big"), int.from_bytes(data[offset + 5:offset + 7], "big")
            offset += 2 + max(block, 0)
    return 160, 120


def _format_datetime(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return text
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(_CHINA).strftime("%Y/%m/%d %H:%M")
